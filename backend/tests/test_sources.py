import hashlib
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

from backend.tests import isolation  # Imported first: keeps the tests off DayWright's own data.
from backend.app import database, sources
from backend.app.sources import (SourceError, briefing_of, file_outline, lesson_address, markdown_outline, parse_page,
                                 scan_folder, section_text, study_minutes, topic_profile)

LESSON = """---
tags: [angular]
---
# Consuming HTTP Services

Angular's HttpClient sends requests to a server and turns the answers into typed data your components can use.
It handles headers, errors and retries in one place.

## HttpClient setup

Provide it once with provideHttpClient().

### Providers

Where the provider goes.

## Interceptors

```ts
export const auth: HttpInterceptorFn = (req, next) => next(req);
```

## Error handling

Exercise: make a request fail and show a message.

```
# not a heading, inside code
```
"""


def tree(root, files):
    """Write `files` ({relative path: text}) under `root`; return root."""
    for relative, text in files.items():
        path = root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
    return root


def hashes(root):
    """Every file under `root` with its content hash and modified time, to show nothing changed."""
    return {str(path.relative_to(root)): (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
            for path in sorted(root.rglob("*")) if path.is_file()}


class FolderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = tree(Path(self.temp.name), {
            "Angular/03 Consuming HTTP Services.md": LESSON, "Angular/01 Basics.md": "# Basics\n\nThe very start.\n",
            "README.md": "# Knowledge\n", ".obsidian/workspace.md": "x", "history/2026-09.md": "# Old\n",
            "Observatory/node_modules/pkg/readme.md": "x", "Observatory/dist/index.md": "x", "image.png": "not text",
            "Notes/.hidden.md": "x"})

    def test_the_tree_lists_readable_files_and_leaves_hidden_history_node_modules_and_dist_unticked(self):
        listed = {item["path"]: item["ticked"] for item in scan_folder(self.root)}

        self.assertEqual(listed, {"Angular/01 Basics.md": True, "Angular/03 Consuming HTTP Services.md": True,
                                  "README.md": True, ".obsidian/workspace.md": False, "history/2026-09.md": False,
                                  "Notes/.hidden.md": False, "Observatory/dist/index.md": False,
                                  "Observatory/node_modules/pkg/readme.md": False})

    def test_a_file_too_large_to_read_is_listed_unticked_with_its_reason(self):
        (self.root / "Angular" / "huge.md").write_bytes(b"#" * 64)
        with patch.object(sources, "FOLDER_MAX_BYTES", 32):
            huge = next(item for item in scan_folder(self.root) if item["path"] == "Angular/huge.md")

        self.assertEqual((huge["ticked"], huge["reason"]), (False, "too large"))

    def test_a_pdf_with_too_many_pages_is_listed_unticked_with_its_reason(self):
        for name, pages in (("long.pdf", 4), ("short.pdf", 3)):
            writer = PdfWriter()
            for _ in range(pages):
                writer.add_blank_page(width=200, height=200)
            with (self.root / name).open("wb") as handle:
                writer.write(handle)
        with patch.object(sources, "FOLDER_MAX_PAGES", 3):
            listed = {item["path"]: (item["ticked"], item["reason"]) for item in scan_folder(self.root) if item["path"].endswith(".pdf")}

        self.assertEqual(listed, {"long.pdf": (False, "too large"), "short.pdf": (True, None)})

    def test_nothing_in_the_folder_changes_whatever_is_read(self):
        before = hashes(self.root)
        scan_folder(self.root)
        for item in scan_folder(self.root):
            file_outline(self.root / item["path"])
        sources.file_hash(self.root / "README.md")

        self.assertEqual(hashes(self.root), before)

    def test_a_file_outside_the_folder_or_a_link_out_of_it_is_never_read(self):
        outside = tree(Path(self.temp.name).parent / f"{Path(self.temp.name).name}-outside", {"secret.md": "# Secret\n"})
        self.addCleanup(lambda: (outside / "secret.md").unlink())
        (self.root / "link.md").symlink_to(outside / "secret.md")

        self.assertNotIn("link.md", [item["path"] for item in scan_folder(self.root)])
        with self.assertRaises(SourceError):
            sources.inside(self.root, "../" + outside.name + "/secret.md")


class OutlineTests(unittest.TestCase):
    def test_markdown_lists_its_first_level_headings_each_with_its_second_level_ones_ignoring_code(self):
        self.assertEqual([(item["title"], [topic["title"] for topic in item["topics"]]) for item in markdown_outline(LESSON)],
                         [("Consuming HTTP Services", ["HttpClient setup", "Interceptors", "Error handling"])])

    def test_second_level_headings_before_any_first_level_one_go_under_the_files_title(self):
        outline = markdown_outline("## Setup\n\n## Use\n", title="Basics")
        self.assertEqual([(item["title"], [topic["title"] for topic in item["topics"]]) for item in outline],
                         [("Basics", ["Setup", "Use"])])

    def test_word_and_pdf_outlines_come_from_their_heading_styles_and_bookmarks(self):
        with tempfile.TemporaryDirectory() as folder:
            word = Path(folder) / "guide.docx"
            document = Document()
            document.add_heading("Routing", level=1)
            document.add_paragraph("Routes map addresses to screens.")
            document.add_heading("Guards", level=2)
            document.add_heading("Lazy loading", level=2)
            document.save(word)
            pdf = Path(folder) / "book.pdf"
            writer = PdfWriter()
            for _ in range(3):
                writer.add_blank_page(200, 200)
            chapter = writer.add_outline_item("Signals", 0)
            writer.add_outline_item("Computed", 1, parent=chapter)
            writer.add_outline_item("Effects", 2, parent=chapter)
            with pdf.open("wb") as handle:
                writer.write(handle)

            outlines = [file_outline(word), file_outline(pdf)]

        self.assertEqual([[(item["title"], [topic["title"] for topic in item["topics"]]) for item in outline] for outline in outlines],
                         [[("Routing", ["Guards", "Lazy loading"])], [("Signals", ["Computed", "Effects"])]])

    def test_a_briefing_is_the_first_paragraph_to_about_forty_words(self):
        self.assertEqual(briefing_of(LESSON), "Angular's HttpClient sends requests to a server and turns the answers into typed data "
                                              "your components can use. It handles headers, errors and retries in one place.")
        long = "# T\n\n" + " ".join(f"word{index}" for index in range(60)) + "\n"
        self.assertEqual(briefing_of(long), " ".join(f"word{index}" for index in range(40)) + "…")
        self.assertIsNone(briefing_of("# Only a heading\n\n## And another\n"))


class WebsiteTests(unittest.TestCase):
    PAGE = """<html><head><title>Observatory · Learning Atlas</title>
    <meta name="description" content="A map of what to learn, lesson by lesson.">
    <script>var hidden = "<h1>Not a heading</h1>";</script></head>
    <body><h1>Angular</h1><p>Long text that is never kept.</p><h2>Signals</h2><h2>Routing <em>basics</em></h2>
    <h1>Python</h1><h2>Typing</h2></body></html>"""

    def test_a_page_keeps_its_title_its_own_briefing_and_its_headings_and_none_of_its_text(self):
        page = parse_page(self.PAGE)

        self.assertEqual(page, {"title": "Observatory · Learning Atlas", "briefing": "A map of what to learn, lesson by lesson.",
                                "outline": [{"title": "Angular", "topics": [{"title": "Signals"}, {"title": "Routing basics"}]},
                                            {"title": "Python", "topics": [{"title": "Typing"}]}]})
        self.assertNotIn("Long text", repr(page))

    def test_the_open_graph_description_stands_in_and_a_bare_page_has_no_briefing(self):
        self.assertEqual(parse_page('<meta property="og:description" content="From og.">')["briefing"], "From og.")
        self.assertEqual(parse_page("<title>App</title><div id=root></div>"),
                         {"title": "App", "briefing": None, "outline": []})

    def test_a_page_is_fetched_once_over_http_with_a_time_limit_and_a_size_cap(self):
        class Answer:
            headers = {"Content-Type": "text/html; charset=utf-8"}

            def read(self, size):
                self.asked = size
                return WebsiteTests.PAGE.encode()

            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

        with patch.object(sources, "urlopen", return_value=Answer()) as opened:
            page = sources.fetch_page("https://observatory.example/")
        self.assertEqual(opened.call_count, 1)
        self.assertEqual(opened.call_args.kwargs["timeout"], sources.FETCH_TIMEOUT)
        self.assertEqual(page["title"], "Observatory · Learning Atlas")
        for address in ("file:///etc/passwd", "ftp://example.com/", "javascript:alert(1)", "observatory.example"):
            with self.subTest(address), self.assertRaises(SourceError):
                sources.fetch_page(address)


class AddressTests(unittest.TestCase):
    SITE = "https://observatory.learning-atlas.workers.dev"

    def test_a_knowledge_transfer_lesson_opens_on_its_own_observatory_page(self):
        self.assertEqual(lesson_address(self.SITE, "AI Expert/01 Foundations & Tools.md"),
                         f"{self.SITE}/lesson/ai-expert-01-foundations-and-tools")
        self.assertEqual(lesson_address(self.SITE + "/", "Git Expert.md"), f"{self.SITE}/lesson/git-expert")

    def test_a_file_that_is_not_a_lesson_opens_its_collection_or_the_site(self):
        self.assertEqual(lesson_address(self.SITE, "AI Expert/slides.pdf"), f"{self.SITE}/library/ai-expert")
        self.assertEqual(lesson_address(self.SITE, "slides.pdf"), self.SITE)

    def test_another_website_opens_at_its_main_address(self):
        self.assertEqual(lesson_address("https://example.com/docs/", "Guide/intro.md"), "https://example.com/docs/")


class ProfileTests(unittest.TestCase):
    def test_a_topic_is_read_from_its_own_section_with_its_subheadings(self):
        setup = section_text(LESSON, "HttpClient setup")
        profile = topic_profile(setup)

        self.assertIn("provideHttpClient", setup)
        self.assertNotIn("Interceptors", setup)
        self.assertEqual(profile["subheadings"], ["Providers"])
        self.assertEqual((profile["handsOn"], profile["effort"]), (False, "light"))

    def test_effort_comes_from_the_text_the_same_way_every_time(self):
        prose = lambda count: " ".join(["word"] * count)
        code = "```py\nprint(1)\n```\n"
        self.assertEqual(topic_profile(prose(200))["effort"], "light")
        self.assertEqual(topic_profile(prose(500))["effort"], "steady")
        self.assertEqual(topic_profile(prose(200) + "\n" + code)["effort"], "steady", "hands-on is never light")
        self.assertEqual(topic_profile(prose(1300))["effort"], "deep")
        self.assertEqual(topic_profile(prose(200) + ("\n" + code) * 3)["effort"], "deep")
        self.assertTrue(topic_profile("Exercise: build it yourself.")["handsOn"])
        self.assertEqual(topic_profile(prose(777)), topic_profile(prose(777)))

    def test_a_study_session_is_estimated_from_the_profile_never_under_thirty_minutes(self):
        self.assertEqual(study_minutes({"words": 100, "codeBlocks": 0, "effort": "light"}), 30)
        self.assertEqual(study_minutes({"words": 1800, "codeBlocks": 0, "effort": "deep"}), 90)
        self.assertEqual(study_minutes({"words": 600, "codeBlocks": 2, "effort": "steady"}), 60)
        self.assertEqual(study_minutes({"effort": "steady"}), 45, "a website's topic, with no text, goes by its effort")

    def test_the_shortest_session_is_the_shortest_task(self):
        self.assertEqual(sources.MIN_STUDY_MINUTES, database.MIN_TASK_MINUTES)


class NoFolderAssumedTests(unittest.TestCase):
    def test_no_folder_or_path_of_this_mac_is_written_into_the_service_or_the_interface(self):
        project = Path(__file__).resolve().parents[2]
        code = [*project.glob("backend/app/*.py"), *project.glob("src/**/*.js"), *project.glob("src/**/*.jsx"),
                project / "src" / "guide" / "guide.json"]
        for path in code:
            text = path.read_text(encoding="utf-8")
            for assumed in ("/Users/", "~/Documents", "Knowledge Transfer", "Professional Quality", "/Volumes/"):
                with self.subTest(path=path.name, assumed=assumed):
                    self.assertNotIn(assumed, text)


if __name__ == "__main__":
    unittest.main()
