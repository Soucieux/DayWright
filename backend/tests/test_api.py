import base64
from contextlib import nullcontext
import sqlite3
import subprocess
import tempfile
import threading
import unittest
from dataclasses import replace
from io import BytesIO
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.app.main import create_app
from backend.app.agents import AgentOrchestrator
from backend.app.database import Database
from backend.app.config import load_settings
from backend.app.local_import import extract_local_file
from backend.app.model_gateway import ModelGateway
from backend.app.retrieval import EMBEDDING_DIMENSION, EmbeddingGateway
from backend.app.speech import SpeechGateway


class FakeGateway:
    def __init__(self):
        self.calls = []

    def status(self):
        return {
            "state": "ready",
            "running": True,
            "runtimeAvailable": True,
            "chatModelAvailable": True,
            "embeddingModelAvailable": True,
            "voiceModelAvailable": True,
            "label": "Test model",
        }

    def reply(self, message, context, system_prompt=None):
        self.calls.append({"message": message, "context": context, "systemPrompt": system_prompt})
        return (f"Considered: {message}", "test-model")

    def stop(self):
        return None


class FakeEmbeddingGateway:
    def __init__(self):
        self.document_calls = []
        self.query_calls = []

    def status(self):
        return {
            "state": "ready",
            "running": True,
            "modelAvailable": True,
            "runtimeAvailable": True,
            "label": "Test embedding model",
            "dimensions": EMBEDDING_DIMENSION,
        }

    def _vector(self, text):
        vector = [0.0] * EMBEDDING_DIMENSION
        lowered = text.lower()
        vector[0 if "sleep" in lowered else 1 if "budget" in lowered else 2] = 1.0
        return vector

    def embed_documents(self, texts):
        self.document_calls.append(texts)
        return [self._vector(text) for text in texts]

    def embed_query(self, text):
        self.query_calls.append(text)
        return self._vector(text)

    def stop(self):
        return None


class FakePublicFetcher:
    def __init__(self):
        self.topics = []

    def fetch(self, topic):
        self.topics.append(topic)
        return {"title": f"Wikipedia · {topic}", "text": f"Public introduction to {topic}.",
                "sourceUrl": "https://en.wikipedia.org/wiki/Public_topic",
                "sourceLicense": "Wikipedia contributors · CC BY-SA 4.0"}


class FakeSpeechGateway:
    def __init__(self):
        self.clips = []

    def status(self):
        return {"state": "available", "modelAvailable": True,
                "runtimeAvailable": True, "label": "Test local Whisper"}

    def transcribe(self, audio):
        self.clips.append(audio)
        return {"text": "Plan a shorter review", "model": "Test local Whisper"}


class SpeechRuntimeTests(unittest.TestCase):
    def test_partial_converted_model_does_not_enable_microphone(self):
        with tempfile.TemporaryDirectory() as folder:
            model_path = Path(folder)
            (model_path / "config.json").write_text("{}")
            (model_path / "tokenizer.json").write_text("{}")
            (model_path / "vocabulary.txt").write_text("test")
            (model_path / "model.bin").write_bytes(b"incomplete")
            settings = replace(load_settings(), whisper_path=model_path)
            status = SpeechGateway(settings).status()
            self.assertEqual(status["state"], "unavailable")
            self.assertFalse(status["modelAvailable"])


class LocalModelRuntimeTests(unittest.TestCase):
    def assert_private_runtime_launch(self, gateway_class):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            binary = root / "llama-server"
            chat_model = root / "chat.gguf"
            embedding_model = root / "embedding.gguf"
            for path in (binary, chat_model, embedding_model):
                path.write_bytes(b"test")
            settings = replace(
                load_settings(),
                database_path=root / "daywright.sqlite3",
                llama_binary=binary,
                model_path=chat_model,
                embedding_path=embedding_model,
            )
            gateway = gateway_class(settings)
            with (
                patch("backend.app.llama_runtime.secrets.token_urlsafe", return_value="private-token"),
                patch("backend.app.llama_runtime.subprocess.Popen") as popen,
                patch("backend.app.llama_runtime.urlopen") as health,
            ):
                popen.return_value.poll.return_value = None
                health.return_value.__enter__.return_value = object()
                self.assertTrue(gateway.start(timeout=0.1)["running"])

            command = popen.call_args.args[0]
            options = popen.call_args.kwargs
            self.assertNotIn("--api-key", command)
            self.assertNotIn("private-token", command)
            self.assertIn("--log-disable", command)
            self.assertEqual(options["env"]["LLAMA_API_KEY"], "private-token")
            self.assertIs(options["stdout"], subprocess.DEVNULL)
            self.assertIs(options["stderr"], subprocess.DEVNULL)
            gateway.stop()

    def test_chat_runtime_keeps_token_out_of_arguments_and_disables_logs(self):
        self.assert_private_runtime_launch(ModelGateway)

    def test_embedding_runtime_keeps_token_out_of_arguments_and_disables_logs(self):
        self.assert_private_runtime_launch(EmbeddingGateway)

    def test_concurrent_starts_wait_for_one_healthy_runtime(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            binary = root / "llama-server"
            model = root / "chat.gguf"
            binary.write_bytes(b"test")
            model.write_bytes(b"test")
            gateway = ModelGateway(replace(
                load_settings(), llama_binary=binary, model_path=model
            ))
            health_started = threading.Event()
            release_health = threading.Event()
            second_started = threading.Event()
            second_finished = threading.Event()
            results = []

            def health(*_args, **_kwargs):
                health_started.set()
                self.assertTrue(release_health.wait(1))
                return nullcontext(object())

            def start(second=False):
                if second:
                    second_started.set()
                results.append(gateway.start(timeout=1)["running"])
                if second:
                    second_finished.set()

            with (
                patch("backend.app.llama_runtime.subprocess.Popen") as popen,
                patch("backend.app.llama_runtime.urlopen", side_effect=health),
            ):
                popen.return_value.poll.return_value = None
                first = threading.Thread(target=start)
                first.start()
                self.assertTrue(health_started.wait(1))
                second = threading.Thread(target=start, kwargs={"second": True})
                second.start()
                self.assertTrue(second_started.wait(1))
                self.assertFalse(second_finished.wait(0.05))
                release_health.set()
                first.join(1)
                second.join(1)

            self.assertEqual(results, [True, True])
            self.assertEqual(popen.call_count, 1)
            gateway.stop()

    def test_spawn_failure_stays_available_without_raising(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            binary = root / "llama-server"
            model = root / "chat.gguf"
            binary.write_bytes(b"test")
            model.write_bytes(b"test")
            gateway = ModelGateway(replace(
                load_settings(), llama_binary=binary, model_path=model
            ))
            with patch(
                "backend.app.llama_runtime.subprocess.Popen",
                side_effect=OSError("cannot execute"),
            ):
                status = gateway.start(timeout=0.1)
            self.assertEqual(status["state"], "available")
            self.assertFalse(status["running"])


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        test_path = Path(self.temp_dir.name) / "test.sqlite3"
        self.gateway = FakeGateway()
        self.embedding_gateway = FakeEmbeddingGateway()
        self.public_fetcher = FakePublicFetcher()
        self.speech_gateway = FakeSpeechGateway()
        app = create_app(
            database_path=test_path,
            gateway=self.gateway,
            embedding_gateway=self.embedding_gateway,
            public_fetcher=self.public_fetcher,
            speech_gateway=self.speech_gateway,
        )
        self.client = TestClient(app)
        self.today = date.today().isoformat()
        self.month = self.today[:7]
        self.yesterday = (date.today() - timedelta(days=1)).isoformat()
        # Explicit isolated prototype fixture for legacy-plan tests, never a runtime default.
        fixture = Database(test_path)
        with fixture.connect() as connection:
            fixture._seed_day(connection, self.today)
        self.day = self.client.get("/api/bootstrap", params={"date": self.today}).json()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_voice_receives_only_bounded_local_wav_and_returns_editable_text(self):
        self.assertEqual(self.day["voice"]["state"], "available")
        refused = self.client.post("/api/voice/transcribe", content=b"private audio",
                                   headers={"content-type": "audio/webm"})
        self.assertEqual(refused.status_code, 415)
        large = self.client.post("/api/voice/transcribe", content=b"A" * 2_000_001,
                                 headers={"content-type": "audio/wav"})
        self.assertEqual(large.status_code, 413)
        clip = self.client.post("/api/voice/transcribe", content=b"short wav clip",
                                headers={"content-type": "audio/wav"})
        self.assertEqual(clip.status_code, 200)
        self.assertEqual(clip.json()["text"], "Plan a shorter review")
        self.assertEqual(self.speech_gateway.clips, [b"short wav clip"])

    def test_knowledge_topic_is_local_first_and_public_fetch_is_topic_only(self):
        self.client.post("/api/knowledge/sources", json={
            "title": "Sleep notebook", "sourceType": "note", "text": "Sleep routine notes from me."})
        local = self.client.post("/api/knowledge/topic", json={"topic": "sleep"}).json()
        self.assertEqual(local["publicFetch"], "not_requested")
        self.assertIsNone(local["source"])
        self.assertEqual(self.public_fetcher.topics, [])

        missing = self.client.post("/api/knowledge/topic", json={"topic": "budget"}).json()
        self.assertEqual(missing["publicFetch"], "awaiting_import_choice")
        self.assertEqual(self.public_fetcher.topics, ["budget"])
        self.assertIsNone(missing["source"])
        self.assertEqual(missing["retrieval"]["status"], "no_match")
        self.assertEqual(missing["importOptions"]["filter"]["priority"],
                         ["credibility", "timeliness", "other"])
        self.assertEqual(len(missing["importOptions"]["plans"]), 3)
        self.assertEqual(self.client.get("/api/knowledge").json()["rag"]["vectorStore"]["sourceCount"], 1)
        pending = self.client.get("/api/knowledge/import-plans").json()["pending"]
        self.assertEqual(pending[0]["topic"], "budget")
        plan_id = missing["importOptions"]["plans"][0]["id"]
        confirmed = self.client.post(f"/api/knowledge/import-plans/{plan_id}/confirm").json()
        self.assertEqual(confirmed["source"]["sourceUrl"],
                         "https://en.wikipedia.org/wiki/Public_topic")
        self.assertEqual(confirmed["organizationLabels"], ["Basic", "Advanced", "Practical"])
        self.assertTrue(self.client.post(f"/api/knowledge/import-plans/{plan_id}/confirm").json()["alreadyImported"])
        other_id = missing["importOptions"]["plans"][1]["id"]
        self.assertEqual(self.client.post(f"/api/knowledge/import-plans/{other_id}/confirm").status_code, 409)
        self.assertEqual(self.client.post("/api/knowledge/search", json={
            "query": "budget"}).json()["status"], "ready")
        chat = self.client.post("/api/chat", json={
            "date": self.today, "message": "budget", "mode": "ask"}).json()
        self.assertEqual(chat["retrieval"]["matches"][0]["sourceUrl"],
                         "https://en.wikipedia.org/wiki/Public_topic")
        history = self.client.get("/api/bootstrap", params={"date": self.today}).json()
        stored_answer = next(turn for turn in history["messages"] if turn["role"] == "assistant")
        self.assertEqual(stored_answer["retrieval"]["matches"][0]["sourceLicense"],
                         "Wikipedia contributors · CC BY-SA 4.0")
        forced = self.client.post("/api/knowledge/topic", json={
            "topic": "sleep", "explicitWeb": True}).json()
        self.assertEqual(forced["publicFetch"], "awaiting_import_choice")
        self.assertEqual(self.public_fetcher.topics, ["budget", "sleep"])
        sources = self.client.get("/api/knowledge").json()["sources"]
        self.assertTrue(any(source["sourceUrl"] for source in sources))
        private = self.client.post("/api/knowledge/topic", json={
            "topic": "my budget account 123456789", "explicitWeb": True}).json()
        self.assertEqual(private["publicFetch"], "needs_general_topic")
        self.assertEqual(self.public_fetcher.topics, ["budget", "sleep"])
        with sqlite3.connect(Path(self.temp_dir.name) / "test.sqlite3") as connection:
            self.assertEqual(connection.execute("PRAGMA quick_check").fetchone()[0], "ok")
            checkpoints_in_main = connection.execute(
                "SELECT COUNT(*) FROM sqlite_master WHERE name = 'checkpoints'"
            ).fetchone()[0]
            self.assertEqual(checkpoints_in_main, 0)
        checkpoint_path = Path(self.temp_dir.name) / "test.checkpoints.sqlite3"
        with sqlite3.connect(checkpoint_path) as connection:
            self.assertEqual(connection.execute("PRAGMA quick_check").fetchone()[0], "ok")
            self.assertGreater(connection.execute("SELECT COUNT(*) FROM checkpoints").fetchone()[0], 0)

    def test_local_file_import_supports_markdown_and_word_but_rejects_unsupported_or_scanned(self):
        filename = base64.b64encode("learn.md".encode()).decode()
        source = self.client.post("/api/knowledge/import", content=b"# Learning\nFrench grammar notes",
                                  headers={"Content-Type": "application/octet-stream",
                                           "X-DayWright-Filename": filename})
        self.assertEqual(source.status_code, 200)
        self.assertEqual(source.json()["sourceType"], "document")
        self.assertTrue(source.json()["title"].startswith("learn.md · "))
        self.assertEqual(self.client.post("/api/knowledge/search", json={
            "query": "French grammar"}).json()["status"], "ready")
        revised = self.client.post("/api/knowledge/import", content=b"# Learning\nFrench review v2",
                                   headers={"Content-Type": "application/octet-stream",
                                            "X-DayWright-Filename": filename})
        self.assertEqual(revised.status_code, 200)
        self.assertNotEqual(revised.json()["id"], source.json()["id"])

        word = Document()
        word.add_paragraph("Learning reflections")
        word.add_table(rows=1, cols=1).cell(0, 0).text = "Keep review blocks short"
        buffer = BytesIO()
        word.save(buffer)
        title, text = extract_local_file("review.docx", buffer.getvalue())
        self.assertEqual(title, "review.docx")
        self.assertIn("Keep review blocks short", text)

        writer = PdfWriter()
        page = writer.add_blank_page(width=100, height=100)
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 70 Td (French lesson) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({
                NameObject("/F1"): DictionaryObject({
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }),
            }),
        })
        output = BytesIO()
        writer.write(output)
        self.assertIn("French lesson", extract_local_file("lesson.pdf", output.getvalue())[1])
        blank = PdfWriter()
        blank.add_blank_page(width=100, height=100)
        output = BytesIO()
        blank.write(output)
        with self.assertRaisesRegex(ValueError, "No extractable text"):
            extract_local_file("scanned.pdf", output.getvalue())
        unsupported = self.client.post("/api/knowledge/import", content=b"not a document",
                                       headers={"Content-Type": "application/octet-stream",
                                                "X-DayWright-Filename": base64.b64encode(
                                                    b"old-word.doc").decode()})
        self.assertEqual(unsupported.status_code, 422)
        self.assertEqual(self.client.post("/api/knowledge/import", content=b"x" * 2_000_001,
                                          headers={"Content-Type": "application/octet-stream",
                                                   "X-DayWright-Filename": filename}).status_code, 413)
        self.assertEqual(self.client.get("/api/knowledge").json()["rag"]["vectorStore"]["sourceCount"], 2)

    def test_bootstrap_returns_three_persisted_variants(self):
        self.assertEqual(len(self.day["variants"]), 3)
        self.assertGreater(len(self.day["entries"]), 5)
        self.assertIsNone(self.day["confirmedVariantId"])

    def test_bootstrap_discloses_bounded_multi_agent_contract(self):
        agents = {agent["key"]: agent for agent in self.day["agents"]}
        self.assertEqual(set(agents), {"orchestrator", "learning", "life", "finance", "summary"})
        self.assertTrue(agents["orchestrator"]["mayProposePlan"])
        self.assertTrue(
            all(not agent["mayProposePlan"] for key, agent in agents.items() if key != "orchestrator")
        )

    def test_confirming_a_variant_is_idempotent_for_the_date(self):
        variant = self.day["variants"][1]
        response = self.client.post(
            "/api/plan/confirm",
            json={"date": self.day["date"], "variantId": variant["id"]},
        )
        self.assertEqual(response.status_code, 200)
        refreshed = self.client.get("/api/bootstrap", params={"date": self.day["date"]}).json()
        self.assertEqual(refreshed["confirmedVariantId"], variant["id"])

    def test_calendar_browses_recorded_days_without_inventing_past_plans(self):
        month = self.client.get("/api/calendar", params={"month": self.month}).json()
        self.assertEqual([day["date"] for day in month["days"]], [self.today])
        empty = self.client.get(
            "/api/bootstrap",
            params={"date": self.yesterday, "create_if_missing": "false"},
        ).json()
        self.assertIsNone(empty["planSetId"])
        self.assertEqual(empty["entries"], [])
        unchanged = self.client.get("/api/calendar", params={"month": self.month}).json()
        self.assertEqual(unchanged["days"], month["days"])

    def test_replacing_a_confirmed_plan_requires_explicit_approval(self):
        balanced, focused = self.day["variants"][:2]
        first = self.client.post(
            "/api/plan/confirm",
            json={"date": self.day["date"], "variantId": balanced["id"]},
        )
        self.assertEqual(first.status_code, 200)
        refused = self.client.post(
            "/api/plan/confirm",
            json={"date": self.day["date"], "variantId": focused["id"]},
        )
        self.assertEqual(refused.status_code, 409)
        current = self.client.get("/api/bootstrap", params={"date": self.day["date"]}).json()
        self.assertEqual(current["confirmedVariantId"], balanced["id"])
        self.assertEqual(current["selectedVariantId"], balanced["id"])
        replaced = self.client.post(
            "/api/plan/confirm",
            json={"date": self.day["date"], "variantId": focused["id"], "replaceExisting": True},
        )
        self.assertEqual(replaced.status_code, 200)
        current = self.client.get("/api/bootstrap", params={"date": self.day["date"]}).json()
        self.assertEqual(current["confirmedVariantId"], focused["id"])
        self.assertEqual(current["selectedVariantId"], focused["id"])
        self.assertEqual(len(current["entries"]), len(self.day["entries"]))

    def test_calendar_summary_and_reporting_follow_current_variant_only(self):
        focused = self.day["variants"][1]
        draft_report = self.client.patch(
            f"/api/entries/{self.day['entries'][0]['id']}", json={"status": "done"},
        )
        self.assertEqual(draft_report.status_code, 409)
        self.client.post(
            "/api/plan/confirm",
            json={"date": self.day["date"], "variantId": focused["id"]},
        )
        focused_day = self.client.get("/api/bootstrap", params={"date": self.day["date"]}).json()
        old_report = self.client.patch(
            f"/api/entries/{self.day['entries'][0]['id']}", json={"status": "done"},
        )
        self.assertEqual(old_report.status_code, 409)
        report = self.client.patch(
            f"/api/entries/{focused_day['entries'][0]['id']}", json={"status": "done"},
        )
        self.assertEqual(report.status_code, 200)
        month = self.client.get("/api/calendar", params={"month": self.month}).json()
        self.assertTrue(month["days"][0]["confirmed"])
        self.assertEqual(month["days"][0]["variantName"], "Focused")
        self.assertEqual(month["days"][0]["doneCount"], 1)

    def test_adjustment_is_a_proposal_until_confirmed(self):
        response = self.client.post(
            "/api/chat",
            json={
                "date": self.day["date"],
                "message": "I am tired; make today gentler",
                "mode": "adjust",
                "selectedVariantId": self.day["selectedVariantId"],
            },
        )
        payload = response.json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(payload["proposedAction"]["payload"]["variantName"], "Gentle")
        self.assertEqual(payload["proposedAction"]["payload"]["proposedBy"], "orchestrator")
        route = [run["agentKey"] for run in payload["agentRoute"]]
        self.assertEqual(route[0], "orchestrator")
        self.assertIn("learning", route)
        self.assertIn("life", route)
        self.assertEqual(route[-1], "orchestrator")
        unchanged = self.client.get("/api/bootstrap", params={"date": self.day["date"]}).json()
        self.assertIsNone(unchanged["confirmedVariantId"])

        decision = self.client.post(
            f"/api/actions/{payload['proposedAction']['id']}",
            json={"decision": "confirmed"},
        )
        self.assertEqual(decision.status_code, 200)
        changed = self.client.get("/api/bootstrap", params={"date": self.day["date"]}).json()
        self.assertEqual(
            changed["confirmedVariantId"], payload["proposedAction"]["payload"]["variantId"]
        )

    def test_chat_plan_replacement_names_current_choice_and_rejects_stale_or_unreviewed_change(self):
        balanced, focused, gentle = self.day["variants"]
        self.client.post("/api/plan/confirm", json={
            "date": self.today, "variantId": balanced["id"]})
        first = self.client.post("/api/chat", json={
            "date": self.today, "message": "I am tired; make today gentler", "mode": "adjust"})
        self.assertEqual(first.status_code, 200)
        proposal = first.json()["proposedAction"]
        self.assertEqual(proposal["payload"]["reviewedFromVariantName"], "Balanced")
        self.assertEqual(proposal["payload"]["reviewedFromVariantId"], balanced["id"])
        self.assertIn("Balanced", proposal["explanation"])
        self.assertIn("Gentle", proposal["explanation"])
        unreviewed = Database(Path(self.temp_dir.name) / "test.sqlite3").propose_action(
            first.json()["threadId"],
            "select_variant", {"date": self.today, "variantId": focused["id"]},
            "Unreviewed plan switch")
        refused = self.client.post(f"/api/actions/{unreviewed['id']}",
                                   json={"decision": "confirmed"})
        self.assertEqual(refused.status_code, 409)
        self.assertEqual(self.client.get("/api/bootstrap", params={"date": self.today}).json()
                         ["confirmedVariantId"], balanced["id"])
        self.client.post("/api/plan/confirm", json={
            "date": self.today, "variantId": focused["id"], "replaceExisting": True})
        stale = self.client.post(f"/api/actions/{proposal['id']}",
                                 json={"decision": "confirmed"})
        self.assertEqual(stale.status_code, 409)
        fresh = self.client.post("/api/chat", json={
            "date": self.today, "message": "I am tired; make today gentler", "mode": "adjust"}).json()
        self.assertEqual(fresh["proposedAction"]["payload"]["reviewedFromVariantName"], "Focused")
        accepted = self.client.post(f"/api/actions/{fresh['proposedAction']['id']}",
                                    json={"decision": "confirmed"})
        self.assertEqual(accepted.status_code, 200)
        self.assertEqual(self.client.get("/api/bootstrap", params={"date": self.today}).json()
                         ["confirmedVariantId"], gentle["id"])

    def test_cross_domain_route_is_visible_and_persisted(self):
        response = self.client.post(
            "/api/chat",
            json={
                "date": self.day["date"],
                "message": "Balance French study, sleep, and my budget",
                "mode": "ask",
                "selectedVariantId": self.day["selectedVariantId"],
            },
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        route = [run["agentKey"] for run in payload["agentRoute"]]
        self.assertEqual(
            route,
            ["orchestrator", "learning", "life", "finance", "summary", "orchestrator"],
        )
        self.assertIn("bounded agent reports", self.gateway.calls[-1]["context"].lower())
        self.assertIn("Orchestrator Agent", self.gateway.calls[-1]["systemPrompt"])

        chinese = self.client.post("/api/chat", json={
            "date": self.day["date"], "message": "请解释今天的计划", "mode": "ask",
            "language": "zh",
        })
        self.assertEqual(chinese.status_code, 200)
        self.assertIn("Respond in Simplified Chinese", self.gateway.calls[-1]["systemPrompt"])

        refreshed = self.client.get(
            "/api/bootstrap", params={"date": self.day["date"]}
        ).json()
        assistant = next(
            message for message in refreshed["messages"] if message["role"] == "assistant"
        )
        self.assertEqual(
            [run["agentKey"] for run in assistant["agentRoute"]], route
        )

    def test_single_domain_request_stays_with_its_agent(self):
        response = self.client.post(
            "/api/chat",
            json={
                "date": self.day["date"],
                "message": "Should I study French longer?",
                "mode": "ask",
                "selectedVariantId": self.day["selectedVariantId"],
            },
        )
        route = [run["agentKey"] for run in response.json()["agentRoute"]]
        self.assertEqual(route, ["orchestrator", "learning", "orchestrator"])

    def test_rag_indexes_retrieves_and_persists_relevant_private_context(self):
        sleep_source = self.client.post(
            "/api/knowledge/sources",
            json={
                "title": "Recovery notes",
                "sourceType": "note",
                "text": "My best sleep routine starts with a screen-free wind-down at 22:30.",
            },
        )
        budget_source = self.client.post(
            "/api/knowledge/sources",
            json={
                "title": "Budget notes",
                "sourceType": "note",
                "text": "Review the grocery budget on Friday before making weekend plans.",
            },
        )
        self.assertEqual(sleep_source.status_code, 200)
        self.assertEqual(budget_source.status_code, 200)

        response = self.client.post(
            "/api/chat",
            json={
                "date": self.day["date"],
                "message": "What did I note about protecting my sleep?",
                "mode": "ask",
                "selectedVariantId": self.day["selectedVariantId"],
            },
        )
        self.assertEqual(response.status_code, 200)
        retrieval = response.json()["retrieval"]
        self.assertEqual(retrieval["status"], "ready")
        self.assertEqual(retrieval["matches"][0]["sourceTitle"], "Recovery notes")
        self.assertIn("screen-free wind-down", self.gateway.calls[-1]["context"])

        knowledge = self.client.get("/api/knowledge").json()
        self.assertEqual(knowledge["rag"]["vectorStore"]["engine"], "sqlite-vec")
        self.assertEqual(knowledge["rag"]["vectorStore"]["sourceCount"], 2)

        refreshed = self.client.get(
            "/api/bootstrap", params={"date": self.day["date"]}
        ).json()
        assistant = next(
            message for message in refreshed["messages"] if message["role"] == "assistant"
        )
        self.assertEqual(
            assistant["retrieval"]["matches"][0]["sourceTitle"], "Recovery notes"
        )

        self.client.post(
            "/api/knowledge/sources",
            json={
                "title": "Recovery notes",
                "sourceType": "note",
                "text": "This revised note replaces the currently indexed source text.",
            },
        )
        history = self.client.get(
            "/api/bootstrap", params={"date": self.day["date"]}
        ).json()
        historical_answer = next(
            message for message in history["messages"] if message["role"] == "assistant"
        )
        self.assertIn(
            "screen-free wind-down",
            historical_answer["retrieval"]["matches"][0]["content"],
        )

    def test_rag_does_not_ground_an_unrelated_question(self):
        self.client.post(
            "/api/knowledge/sources",
            json={
                "title": "Recovery notes",
                "sourceType": "note",
                "text": "My sleep routine starts with a quiet screen-free wind-down.",
            },
        )
        response = self.client.post(
            "/api/knowledge/search",
            json={"query": "What color should I paint a bicycle?", "limit": 4},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "no_match")
        self.assertEqual(response.json()["matches"], [])


class OwnedDayTests(unittest.TestCase):
    """Exercise a genuine empty account, not the opt-in legacy example fixture."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.client = TestClient(create_app(
            database_path=Path(self.temp_dir.name) / "owned.sqlite3",
            gateway=FakeGateway(), embedding_gateway=FakeEmbeddingGateway(),
        ))
        self.today = date.today().isoformat()
        self.yesterday = (date.today() - timedelta(days=1)).isoformat()

    def tearDown(self):
        self.temp_dir.cleanup()

    def item(self, title, start, domain, goal_id=None, protected=False, repeat="none"):
        return {"date": self.today, "title": title, "detail": "User-owned commitment",
                "domain": domain, "startTime": start, "durationMinutes": 60,
                "constraintKind": "flexible", "repeatKind": repeat,
                "protected": protected, "goalId": goal_id}

    def test_an_unused_record_can_be_removed_but_a_confirmed_one_cannot(self):
        keep = self.client.post("/api/daily-items", json=self.item("Dentist", "09:00", "life"))
        mistake = self.client.post("/api/daily-items", json=self.item("Typo task", "11:00", "life"))
        self.assertEqual(mistake.status_code, 200)

        removed = self.client.delete(f"/api/daily-items/{mistake.json()['id']}")
        self.assertEqual(removed.status_code, 200)
        self.assertEqual(removed.json()["title"], "Typo task")
        remaining = self.client.get("/api/bootstrap", params={"date": self.today}).json()["dayItems"]
        self.assertEqual([item["title"] for item in remaining], ["Dentist"])
        self.assertEqual(
            self.client.delete(f"/api/daily-items/{mistake.json()['id']}").status_code, 404
        )

        plan = self.client.post("/api/plan/generate", json={"date": self.today}).json()
        self.assertEqual(self.client.post("/api/plan/confirm", json={
            "date": self.today, "variantId": plan["variants"][0]["id"]}).status_code, 200)
        blocked = self.client.delete(f"/api/daily-items/{keep.json()['id']}")
        self.assertEqual(blocked.status_code, 409)
        still_recorded = self.client.get("/api/bootstrap", params={"date": self.today}).json()
        self.assertEqual([item["title"] for item in still_recorded["dayItems"]], ["Dentist"])

    def test_a_goal_can_be_removed_only_after_its_linked_records(self):
        goal = self.client.post("/api/goals", json={
            "title": "Learn French", "domain": "learning"}).json()
        linked = self.client.post("/api/daily-items", json=self.item(
            "French practice", "09:00", "learning", goal["id"])).json()

        blocked = self.client.delete(f"/api/goals/{goal['id']}")
        self.assertEqual(blocked.status_code, 409)
        self.assertIn("still link", blocked.json()["detail"])

        self.assertEqual(self.client.delete(f"/api/daily-items/{linked['id']}").status_code, 200)
        self.assertEqual(self.client.delete(f"/api/goals/{goal['id']}").status_code, 200)
        self.assertEqual(
            self.client.get("/api/bootstrap", params={"date": self.today}).json()["goals"], []
        )
        self.assertEqual(self.client.delete(f"/api/goals/{goal['id']}").status_code, 404)

    def test_empty_account_and_past_day_are_not_invented(self):
        fresh = self.client.get("/api/bootstrap", params={"date": self.today}).json()
        self.assertIsNone(fresh["planSetId"])
        self.assertEqual(fresh["dayItems"], [])
        self.assertEqual(fresh["goals"], [])
        legacy_flag = self.client.get("/api/bootstrap", params={
            "date": self.today, "create_if_missing": "true"}).json()
        self.assertIsNone(legacy_flag["planSetId"])
        self.assertEqual(self.client.post("/api/plan/generate", json={
            "date": self.today}).status_code, 422)
        self.assertEqual(self.client.get("/api/calendar", params={"month": self.today[:7]}).json()["days"], [])
        past = self.item("Old appointment", "09:00", "life")
        past["date"] = self.yesterday
        self.assertEqual(self.client.post("/api/daily-items", json=past).status_code, 409)
        self.assertEqual(self.client.post("/api/plan/generate", json={"date": self.today}).status_code, 422)

    def test_past_daily_summary_is_frozen_after_first_saved_report(self):
        self.assertEqual(
            self.client.get("/api/summaries", params={"date": self.yesterday}).status_code,
            405,
        )
        earlier = self.client.post("/api/summaries", params={"date": self.yesterday}).json()
        self.assertEqual(earlier["reports"]["day"]["goals"], [])
        self.client.post("/api/goals", json={"title": "New goal today", "domain": "life"})
        refreshed = self.client.post("/api/summaries", params={"date": self.yesterday}).json()
        self.assertEqual(refreshed["reports"]["day"], earlier["reports"]["day"])

    def test_future_commitment_is_editable_but_future_plan_and_outcome_are_not(self):
        future = self.item("Next week's appointment", "10:00", "life")
        future["date"] = (date.today() + timedelta(days=7)).isoformat()
        created = self.client.post("/api/daily-items", json=future)
        self.assertEqual(created.status_code, 200)
        record = created.json()
        future["detail"] = "Confirmed appointment details"
        update = self.client.put(f"/api/daily-items/{record['id']}", json={**future, "status": "planned"})
        self.assertEqual(update.status_code, 200)
        self.assertEqual(update.json()["detail"], "Confirmed appointment details")
        premature = self.client.put(f"/api/daily-items/{record['id']}", json={**future, "status": "done"})
        self.assertEqual(premature.status_code, 409)
        self.assertEqual(self.client.post("/api/plan/generate", json={"date": future["date"]}).status_code, 409)
        self.assertEqual(self.client.get("/api/bootstrap", params={"date": future["date"]}).json()["dayItems"][0]["completion_status"], "planned")

    def test_learning_subject_and_explicit_session_reach_summary_and_learning_agent(self):
        subject = self.client.post("/api/learning/items", json={
            "title": "French pronunciation", "difficulty": "hard",
            "estimatedMinutes": 35})
        self.assertEqual(subject.status_code, 200)
        item_id = subject.json()["id"]
        session = self.client.post("/api/learning/sessions", json={
            "date": self.today, "itemId": item_id, "minutes": 25, "result": "done"})
        self.assertEqual(session.status_code, 200)
        area = self.client.get("/api/areas/learning", params={"date": self.today}).json()
        self.assertEqual(area["items"][0]["difficulty"], "hard")
        self.assertEqual(area["sessions"][0]["minutes"], 25)
        self.assertEqual(self.client.post("/api/learning/sessions", json={
            "date": self.yesterday, "itemId": item_id, "minutes": 25,
            "result": "done"}).status_code, 409)
        summary = self.client.post("/api/summaries", params={"date": self.today}).json()
        self.assertEqual(summary["reports"]["day"]["areaEvidence"]["learning"]["sessions"], 1)
        self.assertEqual(summary["reports"]["day"]["recordedDays"], 1)
        asked = self.client.post("/api/chat", json={
            "date": self.today, "message": "How is French learning?", "mode": "ask"}).json()
        learning = next(run for run in asked["agentRoute"] if run["agentKey"] == "learning")
        self.assertIn("1 explicitly reported session", learning["summary"])
        self.assertEqual(self.client.patch(f"/api/learning/items/{item_id}", json={
            "status": "archived"}).status_code, 200)
        self.assertEqual(self.client.post("/api/learning/sessions", json={
            "date": self.today, "itemId": item_id, "minutes": 20,
            "result": "done"}).status_code, 422)

    def test_summary_advice_names_the_record_and_a_concrete_next_plan_change(self):
        created = self.client.post("/api/daily-items", json=self.item(
            "Review retrieval notes", "09:30", "learning", protected=True))
        self.assertEqual(created.status_code, 200)
        plan = self.client.post("/api/plan/generate", json={"date": self.today}).json()
        selected = plan["variants"][0]["id"]
        self.assertEqual(self.client.post("/api/plan/confirm", json={
            "date": self.today, "variantId": selected}).status_code, 200)
        entry = next(item for item in plan["entries"] if item["title"] == "Review retrieval notes")
        self.assertEqual(self.client.patch(f"/api/entries/{entry['id']}", json={
            "status": "skipped"}).status_code, 200)

        report = self.client.post("/api/summaries", params={"date": self.today}).json()
        advice = next(item for item in report["reports"]["day"]["suggestions"]
                      if "Review retrieval notes" in item["content"])
        self.assertIn("45-minute version at 09:30", advice["content"])
        self.assertIn("User-owned commitment", advice["content"])
        self.assertNotIn("Review the size or timing", advice["content"])

    def test_summary_pool_retires_superseded_active_advice_for_the_same_period(self):
        store = Database(Path(self.temp_dir.name) / "pool.sqlite3")
        store.sync_suggestion_pool({"periodKind": "day", "periodKey": self.today,
                                    "suggestions": [{"domain": "life", "priority": "soft",
                                                     "content": "Generic old advice"}]})
        store.sync_suggestion_pool({"periodKind": "day", "periodKey": self.today,
                                    "suggestions": [{"domain": "life", "priority": "soft",
                                                     "content": "Specific current advice"}]})
        items = store.suggestion_pool("day", self.today)["items"]
        self.assertEqual([item["content"] for item in items], ["Specific current advice"])

    def test_life_state_habit_and_event_link_calendar_and_summary_to_next_plan(self):
        habit = self.client.post("/api/life/habits", json={
            "title": "Evening walk", "frequency": "daily"}).json()
        self.assertEqual(self.client.put(
            f"/api/life/habits/{habit['id']}/logs/{self.today}",
            json={"done": True, "note": "Restored energy"}).status_code, 200)
        daily = self.client.put(f"/api/life/daily/{self.today}", json={
            "sleepHours": 6.5, "energyLevel": 2, "mood": 3,
            "note": "Tired after errands"})
        self.assertEqual(daily.status_code, 200)
        event = self.client.post("/api/life/events", json={
            "date": self.today, "title": "Gym appointment", "startTime": "17:30",
            "endTime": "18:15", "category": "sport", "flexible": False})
        self.assertEqual(event.status_code, 200)
        self.assertEqual(self.client.post("/api/life/events", json={
            "date": self.today, "title": "Overlap", "startTime": "18:00",
            "endTime": "18:30", "category": "social"}).status_code, 422)
        self.assertEqual(self.client.post("/api/life/events", json={
            "date": self.yesterday, "title": "Old event", "startTime": "09:00",
            "endTime": "10:00"}).status_code, 409)
        self.assertEqual(self.client.put(f"/api/life/daily/{self.yesterday}", json={
            "energyLevel": 2}).status_code, 409)
        life = self.client.get("/api/areas/life", params={"date": self.today}).json()
        self.assertEqual(life["events"][0]["itemId"], event.json()["itemId"])
        self.assertEqual(life["daily"]["energyLevel"], 2)
        self.assertEqual(life["logs"][0]["done"], 1)
        misplaced = {**self.item("Gym appointment", "17:30", "finance"),
                     "durationMinutes": 45, "status": "planned"}
        self.assertEqual(self.client.put(
            f"/api/daily-items/{event.json()['itemId']}", json=misplaced).status_code, 422)
        calendar = self.client.get("/api/calendar", params={"month": self.today[:7]}).json()
        self.assertEqual(calendar["days"][0]["managedCount"], 1)
        report = self.client.post("/api/summaries", params={"date": self.today}).json()
        self.assertIn("energy was 2/5", " ".join(
            item["content"] for item in report["reports"]["day"]["suggestions"]))
        self.client.post("/api/daily-items", json=self.item(
            "Flexible care", "10:00", "life"))
        plan = self.client.post("/api/plan/generate", json={"date": self.today}).json()
        gentle = next(variant for variant in plan["variants"] if variant["slug"] == "gentle")
        selected = self.client.get("/api/bootstrap", params={
            "date": self.today, "variant_id": gentle["id"]}).json()
        self.assertEqual(next(entry for entry in selected["entries"]
                              if entry["title"] == "Flexible care")["duration_minutes"], 45)

    def test_money_transactions_budget_and_manual_balance_reach_summary_without_bank_sync(self):
        self.assertEqual(self.client.put("/api/money/opening-balance", json={
            "cents": 10_000}).status_code, 200)
        for kind, cents in (("expense", 1_250), ("income", 4_000)):
            self.assertEqual(self.client.post("/api/money/transactions", json={
                "date": self.today, "type": kind, "amountCents": cents,
                "category": "Groceries", "note": "Manual entry"}).status_code, 200)
        self.assertEqual(self.client.put("/api/money/budgets", json={
            "month": self.today[:7], "category": "Groceries",
            "budgetCents": 1_000}).status_code, 200)
        finance = self.client.get("/api/areas/finance", params={"date": self.today}).json()
        self.assertEqual(finance["balanceCents"], 12_750)
        self.assertEqual(finance["budgets"][0]["spentCents"], 1_250)
        self.assertEqual(finance["budgets"][0]["budgetCents"], 1_000)
        self.assertEqual(self.client.post("/api/money/transactions", json={
            "date": self.yesterday, "type": "expense", "amountCents": 100,
            "category": "Old"}).status_code, 409)
        self.assertEqual(self.client.put("/api/money/budgets", json={
            "month": "2025-01", "category": "Old", "budgetCents": 100}).status_code, 409)
        report = self.client.post("/api/summaries", params={"date": self.today}).json()
        self.assertEqual(report["reports"]["day"]["areaEvidence"]["finance"]["transactions"], 2)
        self.assertIn("exceeded", " ".join(
            item["content"] for item in report["reports"]["day"]["suggestions"]))
        asked = self.client.post("/api/chat", json={
            "date": self.today, "message": "How are my groceries and budget?", "mode": "ask"}).json()
        money_agent = next(run for run in asked["agentRoute"] if run["agentKey"] == "finance")
        self.assertIn("12.50 spent of 10.00", money_agent["summary"])

    def test_owned_records_goals_agent_plan_confirmation_and_summaries(self):
        goal = self.client.post("/api/goals", json={"title": "Learn French", "domain": "learning"})
        self.assertEqual(goal.status_code, 200)
        study = self.client.post("/api/daily-items", json=self.item(
            "French practice", "09:00", "learning", goal.json()["id"], True))
        self.assertEqual(study.status_code, 200)
        chore = self.client.post("/api/daily-items", json=self.item("Housework", "11:00", "life"))
        self.assertEqual(chore.status_code, 200)
        before = self.client.get("/api/bootstrap", params={"date": self.today}).json()
        self.assertIsNone(before["planSetId"])
        self.assertEqual(len(before["dayItems"]), 2)
        plan = self.client.post("/api/plan/generate", json={"date": self.today})
        self.assertEqual(plan.status_code, 200)
        planned = plan.json()
        self.assertEqual(planned["planSource"], "orchestrator-records-v1")
        self.assertEqual({variant["name"] for variant in planned["variants"]}, {"Balanced", "Focused", "Gentle"})
        self.assertEqual([entry["title"] for entry in planned["entries"]], ["French practice", "Housework"])
        self.assertEqual(planned["planRoute"][0]["agentKey"], "orchestrator")
        self.assertEqual(planned["planRoute"][-1]["agentKey"], "orchestrator")
        self.assertIsNone(planned["confirmedVariantId"])
        selected = planned["variants"][0]["id"]
        self.assertEqual(self.client.post("/api/plan/confirm", json={"date": self.today, "variantId": selected}).status_code, 200)
        self.assertEqual(self.client.patch(f"/api/entries/{planned['entries'][0]['id']}", json={"status": "done"}).status_code, 200)
        current = self.client.get("/api/bootstrap", params={"date": self.today}).json()
        self.assertEqual(current["confirmedVariantId"], selected)
        self.assertEqual(next(item for item in current["dayItems"] if item["title"] == "French practice")["completion_status"], "done")
        self.assertEqual(next(row for row in current["goals"] if row["id"] == goal.json()["id"])["doneCount"], 1)
        linked = next(row for row in current["goals"] if row["id"] == goal.json()["id"])["linkedItems"]
        self.assertEqual([(item["title"], item["status"]) for item in linked], [("French practice", "done")])
        month = self.client.get("/api/calendar", params={"month": self.today[:7]}).json()
        self.assertEqual(month["days"][0]["doneCount"], 1)
        reports = self.client.post("/api/summaries", params={"date": self.today}).json()["reports"]
        self.assertEqual(set(reports), {"day", "week", "month"})
        self.assertEqual(reports["day"]["domains"]["learning"]["done"], 1)

    def test_adjustment_request_is_explicit_memory_and_does_not_change_plan(self):
        self.client.post("/api/daily-items", json=self.item("French practice", "09:00", "learning", protected=True))
        self.client.post("/api/plan/generate", json={"date": self.today})
        message = self.client.post("/api/chat", json={"date": self.today,
            "message": "Please shorten French practice today", "mode": "adjust"})
        self.assertEqual(message.status_code, 200)
        again = self.client.post("/api/chat", json={"date": self.today,
            "message": "Please shorten French practice more", "mode": "adjust"})
        self.assertEqual(again.status_code, 200)
        self.assertIsNone(self.client.get("/api/bootstrap", params={"date": self.today}).json()["confirmedVariantId"])
        report = self.client.post("/api/summaries", params={"date": self.today}).json()["reports"]["day"]
        self.assertEqual(report["feedback"][0]["taskTitle"], "French practice")
        self.assertTrue(report["feedback"][0]["protected"])
        self.assertIn("keep", " ".join(s["content"] for s in report["suggestions"]).lower())

    def test_summary_suggestion_pool_persists_discard_and_notices_future_repeat(self):
        self.client.post("/api/daily-items", json=self.item(
            "French practice", "09:00", "learning", protected=True))
        for phrase in ("Please shorten French practice", "Please shorten French practice again"):
            self.assertEqual(self.client.post("/api/chat", json={"date": self.today,
                "message": phrase, "mode": "adjust"}).status_code, 200)
        report = self.client.post("/api/summaries", params={"date": self.today}).json()
        day = next(item for item in report["pool"]["day"]["items"]
                   if "French practice" in item["content"])
        self.assertEqual(day["priority"], "strong")
        self.assertEqual(day["status"], "active")
        self.assertGreaterEqual(len(report["pool"]["week"]["items"]), 1)
        planned = self.client.post("/api/plan/generate", json={"date": self.today}).json()
        learning_run = next(run for run in planned["planRoute"] if run["agentKey"] == "learning")
        self.assertIn("Active Summary guidance: strong", learning_run["summary"])
        discarded = self.client.post(f"/api/suggestion-pool/{day['id']}/discard")
        self.assertEqual(discarded.status_code, 200)
        self.assertGreaterEqual(discarded.json()["affectedPeriods"], 3)
        repeated = self.client.post("/api/summaries", params={"date": self.today}).json()
        self.assertEqual(next(item for item in repeated["pool"]["day"]["items"]
                              if item["id"] == day["id"])["status"], "discarded")
        week = repeated["pool"]["week"]["periodKey"]
        refused = self.client.post("/api/suggestion-pool/clear-week", json={
            "week": week, "domain": "learning", "confirmation": "wrong"})
        self.assertEqual(refused.status_code, 422)
        self.assertGreaterEqual(len(repeated["pool"]["week"]["items"]), 1)
        cleared = self.client.post("/api/suggestion-pool/clear-week", json={
            "week": week, "domain": "learning",
            "confirmation": f"CLEAR {week} LEARNING"})
        self.assertEqual(cleared.status_code, 200)
        self.assertGreaterEqual(cleared.json()["deletedAdvice"], 1)
        self.assertFalse(self.client.post("/api/summaries", params={
            "date": self.today}).json()["pool"]["week"]["items"])
        store = Database(Path(self.temp_dir.name) / "owned.sqlite3")
        self.assertFalse(any("French practice" in item["content"]
                             for item in store.active_suggestion_pool()))
        store.sync_suggestion_pool({"periodKind": "day", "periodKey": "later-day",
                                    "suggestions": [{"domain": "learning",
                                                     "content": day["content"],
                                                     "priority": "strong"}]})
        later = store.suggestion_pool("day", "later-day", "learning")
        self.assertEqual(later["items"], [])
        self.assertEqual(len(later["notices"]), 1)
        self.assertEqual(self.client.post(f"/api/suggestion-pool/{day['id']}/discard").status_code, 404)

    def test_summary_adds_traceable_future_work_and_agent_can_explain_and_shorten(self):
        self.client.post("/api/daily-items", json=self.item(
            "French practice", "09:00", "learning", protected=True, repeat="daily"))
        for phrase in ("Please shorten French practice", "Please shorten French practice again"):
            response = self.client.post("/api/chat", json={"date": self.today,
                "message": phrase, "mode": "adjust"})
            self.assertEqual(response.status_code, 200)
        first = self.client.post("/api/summaries", params={"date": self.today})
        self.assertEqual(first.status_code, 200)
        prepared = first.json()["futurePrepared"]
        self.assertEqual(len(prepared), 1)
        future = prepared[0]
        self.assertEqual(future["date"], (date.today() + timedelta(days=1)).isoformat())
        self.assertEqual(future["originKind"], "agent-origin")
        self.assertEqual(future["duration_minutes"], 45)
        self.assertIn("Summary", future["originDetail"])
        self.assertEqual(self.client.post("/api/summaries", params={"date": self.today}).json()["futurePrepared"], [])
        why = self.client.post("/api/chat", json={"date": future["date"],
            "message": "Why was French practice added?", "mode": "ask"})
        self.assertEqual(why.status_code, 200)
        self.assertIn("Record provenance", why.json()["assistantMessage"]["content"])
        proposal = self.client.post("/api/chat", json={"date": future["date"],
            "message": "Please shorten French practice", "mode": "adjust"})
        self.assertEqual(proposal.status_code, 200)
        action = proposal.json()["proposedAction"]
        self.assertEqual(action["actionType"], "shorten_future_item")
        self.assertEqual(action["payload"]["durationMinutes"], 30)
        untouched = self.client.get("/api/bootstrap", params={"date": future["date"]}).json()
        self.assertEqual(untouched["dayItems"][0]["duration_minutes"], 45)
        self.assertEqual(self.client.post(f"/api/actions/{action['id']}", json={"decision": "confirmed"}).status_code, 200)
        edited = self.client.get("/api/bootstrap", params={"date": future["date"]}).json()
        self.assertEqual(edited["dayItems"][0]["duration_minutes"], 30)
        self.assertEqual(edited["dayItems"][0]["originKind"], "agent-origin")
        self.assertIsNone(edited["planSetId"])
        with patch("backend.app.database._writable_day"):
            arrived = self.client.post("/api/plan/generate", json={"date": future["date"]})
        self.assertEqual(arrived.status_code, 200)
        self.assertEqual(arrived.json()["entries"][0]["title"], "French practice")
        self.assertIsNone(arrived.json()["confirmedVariantId"])

    def test_summary_prepares_repeatedly_completed_protected_work_with_shorten_precedence(self):
        store = Database(Path(self.temp_dir.name) / "test.sqlite3")
        today = date.today()
        first = today - timedelta(days=today.weekday())
        if first == today:
            first -= timedelta(days=7)
        second = first + timedelta(days=1)
        with store.connect() as connection:
            for recorded_date in (first, second):
                for title, domain, time, minutes in (
                    ("French practice", "learning", "09:00", 60),
                    ("Morning walk", "life", "10:00", 30),
                ):
                    connection.execute(
                        """INSERT INTO daily_items
                           (id, item_date, title, domain, start_time, duration_minutes,
                            constraint_kind, repeat_kind, protected, completion_status, created_at)
                           VALUES (?, ?, ?, ?, ?, ?, 'flexible', 'daily', 1, 'done', ?)""",
                        (f"{title}:{recorded_date}", recorded_date.isoformat(), title,
                         domain, time, minutes, recorded_date.isoformat()),
                    )
        facts = store.summary_facts(first.isoformat(), second.isoformat())
        self.assertEqual({entry["taskTitle"] for entry in facts["completedRecurring"]},
                         {"French practice", "Morning walk"})
        orchestrator = AgentOrchestrator()
        report = orchestrator.summary_report("week", "synthetic-week", facts)
        self.assertEqual(len(report["completedRecurring"]), 2)
        report["feedback"].append({"taskTitle": "French practice", "domain": "learning",
                                   "shortenRequests": 2, "protected": True})
        prepared = orchestrator.prepare_future_from_summary(store, report)
        self.assertEqual(len(prepared), 2)
        by_title = {item["title"]: item for item in prepared}
        self.assertEqual(by_title["Morning walk"]["duration_minutes"], 30)
        self.assertIn("completed on 2 recorded days", by_title["Morning walk"]["originDetail"])
        self.assertEqual(by_title["French practice"]["duration_minutes"], 45)
        self.assertIn("requests to shorten", by_title["French practice"]["originDetail"])
        self.assertTrue(all(item["originKind"] == "agent-origin" for item in prepared))
        self.assertEqual(orchestrator.prepare_future_from_summary(store, report), [])


if __name__ == "__main__":
    unittest.main()
