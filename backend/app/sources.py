"""What DayWright reads from a source the user connects: a folder on this Mac, which it reads and never
writes, or a website, which it looks into once and keeps only the title, briefing and headings of.

From either it takes an outline (first-level headings, each with its second-level ones, which become
goals and their topics), a briefing (what the source is about, in about forty words), and, for a
topic in a local file, a profile of its text (its subheadings, size, whether it is hands-on, and how
much effort it takes), worked out the same way every time. No folder or path is assumed: the user
picks any folder, anywhere.
"""

from __future__ import annotations

import hashlib
import os
import re
from html.parser import HTMLParser
from math import ceil
from pathlib import Path, PurePosixPath
from urllib.parse import unquote, urlsplit
from urllib.request import Request, urlopen

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .planner import SLOT_MINUTES
from .wording import count_words, cut_words

# The files a folder's tree lists: Markdown, PDF and Word.
READABLE = (".md", ".markdown", ".pdf", ".docx")
# Folders whose files start unticked, as history, installed packages and build output, besides hidden ones.
SKIPPED_FOLDERS = ("history", "node_modules", "dist")
# The largest file, and the most PDF pages, a folder's file may have to be read.
FOLDER_MAX_BYTES = 20_000_000
FOLDER_MAX_PAGES = 300
# A website is given this many seconds to answer, and this much of its page is read.
FETCH_TIMEOUT = 10
FETCH_MAX_BYTES = 2_000_000
# The Observatory site, whose lessons each have a page of their own when a folder's website is this address.
OBSERVATORY_HOSTS = ("observatory.learning-atlas.workers.dev",)
# Effort: under this many words and not hands-on is light; this many words, or this many code examples, is deep.
LIGHT_WORDS = 300
DEEP_WORDS = 1200
DEEP_CODE_BLOCKS = 3
# A study session's length: the topic's words at this pace, and this long for each code example.
STUDY_WORDS_PER_MINUTE = 20
CODE_BLOCK_MINUTES = 15
# A topic without text, from a website, goes by its effort.
# A website keeps no text, so its page is weighed by its headings: a section takes SECTION_MINUTES, and
# up to LIGHT_SECTIONS sections are light, from DEEP_SECTIONS deep.
SECTION_MINUTES = 15
LIGHT_SECTIONS = 2
DEEP_SECTIONS = 6
# The shortest study session, as every task's length is at least this.
MIN_STUDY_MINUTES = 30

# Words that make a topic hands-on even without code: an exercise to do.
_HANDS_ON = re.compile(r"\b(exercises?|try it|practi[cs]e|lab|hands-on)\b|练习", re.IGNORECASE)
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_FENCE = re.compile(r"^\s*(```|~~~)")
# An Obsidian block anchor at a line's end, such as " ^abc123", which Observatory leaves out of a slug.
_BLOCK_ANCHOR = re.compile(r"(?:^|[ \t])\^[A-Za-z0-9_-]+(?=[ \t]*(?:\n|$))")
_LINK = re.compile(r"!?\[([^\]]*)\]\([^)]*\)")
_MARKUP = re.compile(r"[*_`~]|<[^>]+>")
_WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'’_-]*")
_CJK = re.compile(r"[㐀-鿿]")


class SourceError(ValueError):
    """A source that can't be read or reached, with the reason in words for the user."""


class SourceNotFound(SourceError):
    """A source, folder or topic the Library doesn't have."""


def inside(root: Path, relative: str) -> Path:
    """Return a file of a folder by its path in it, refusing one that leads out of the folder.

    Raises:
        SourceError: When the path, links included, ends outside the folder.
    """
    base = root.resolve()
    path = (base / relative).resolve()
    if base not in path.parents:
        raise SourceError("That file is outside the folder.")
    return path


def scan_folder(root: Path) -> list[dict]:
    """List a folder's readable files for its tree of tick boxes, without changing anything in it.

    Hidden files and folders, and the folders in SKIPPED_FOLDERS, start unticked, as does a file too
    large to read; a link is never followed, so nothing outside the folder is listed.

    Returns:
        Each Markdown, PDF or Word file, by its path in the folder, with its size, whether it starts
        ticked, and why not ("hidden", "skipped" or "too large"), in path order.
    """
    base = root.resolve()
    items = []
    for folder, folders, files in os.walk(base, followlinks=False):
        folders[:] = sorted(name for name in folders if not os.path.islink(os.path.join(folder, name)))
        for name in sorted(files):
            path = Path(folder) / name
            if path.suffix.lower() not in READABLE or path.is_symlink():
                continue
            parts = path.relative_to(base).parts
            size = path.stat().st_size
            reason = ("hidden" if any(part.startswith(".") for part in parts)
                      else "skipped" if any(part in SKIPPED_FOLDERS for part in parts[:-1])
                      else "too large" if size > FOLDER_MAX_BYTES or _too_many_pages(path) else None)
            items.append({"path": PurePosixPath(*parts).as_posix(), "size": size, "ticked": reason is None, "reason": reason})
    return sorted(items, key=lambda item: item["path"])


def file_hash(path: Path) -> str:
    """The SHA-256 of a file's content, which tells a changed file from the one read before."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _too_many_pages(path: Path) -> bool:
    """Whether a PDF has more than FOLDER_MAX_PAGES pages; one that can't be opened says so when it is read."""
    if path.suffix.lower() != ".pdf":
        return False
    try:
        return len(PdfReader(path).pages) > FOLDER_MAX_PAGES
    except (PdfReadError, OSError):
        return False


def _pdf(path: Path) -> PdfReader:
    """Open a PDF to read, refusing a locked one or one with more than FOLDER_MAX_PAGES pages."""
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise SourceError("Unlock the PDF to read it.")
        if len(reader.pages) > FOLDER_MAX_PAGES:
            raise SourceError(f"PDFs of up to {FOLDER_MAX_PAGES} pages can be read.")
    except (PdfReadError, OSError) as error:
        raise SourceError("This PDF could not be read.") from error
    return reader


def read_file(path: Path) -> str:
    """A Markdown, PDF or Word file's text, opened only to read.

    Raises:
        SourceError: When the file is too large, locked, or can't be read.
    """
    if path.stat().st_size > FOLDER_MAX_BYTES:
        raise SourceError("This file is too large to read.")
    suffix = path.suffix.lower()
    try:
        if suffix in (".md", ".markdown"):
            return path.read_bytes().decode("utf-8-sig")
        if suffix == ".pdf":
            return "\n".join(page.extract_text() or "" for page in _pdf(path).pages)
        return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    except SourceError:
        raise
    except (UnicodeError, OSError, ValueError) as error:
        raise SourceError("This file could not be read.") from error


# The revision an import adds to a file's name, which the Library leaves out when it names the file.
_REVISION = re.compile(r" · [0-9a-f]{12}$")


def library_name(title: str) -> str:
    """A source's name as the Library shows it: an imported file's without the revision its import added."""
    return _REVISION.sub("", title)


def plain(text: str) -> str:
    """Text without Markdown or HTML markup, its spaces run together."""
    return " ".join(_MARKUP.sub("", _LINK.sub(r"\1", text)).split())


def _lines(text: str) -> list[tuple[int, str, bool]]:
    """Each line of Markdown with its number and whether it is code (a fence or inside one); front matter is left out."""
    lines = text.splitlines()
    start = 0
    if lines and lines[0].strip() == "---":
        start = next((index + 1 for index in range(1, len(lines)) if lines[index].strip() == "---"), 0)
    marked, fenced = [], False
    for number in range(start, len(lines)):
        line = lines[number]
        if _FENCE.match(line):
            fenced = not fenced
            marked.append((number, line, True))
        else:
            marked.append((number, line, fenced))
    return marked


def _headings(text: str):
    """Each Markdown heading outside code, as (line number, level, plain title)."""
    for number, line, fenced in _lines(text):
        match = None if fenced else _HEADING.match(line)
        if match:
            yield number, len(match.group(1)), plain(match.group(2))


def markdown_outline(text: str, title: str | None = None) -> list[dict]:
    """A Markdown text's first-level headings, each with its second-level ones as topics.

    Args:
        text: The Markdown.
        title: What second-level headings before any first-level one go under, such as the file's name.

    Returns:
        [{"title", "line", "topics": [{"title", "line"}]}], in the order they come.
    """
    outline: list[dict] = []
    for number, level, heading in _headings(text):
        if level == 1:
            outline.append({"title": heading, "line": number, "topics": []})
        elif level == 2:
            if not outline:
                outline.append({"title": title or heading, "line": None, "topics": []})
            outline[-1]["topics"].append({"title": heading, "line": number})
    return outline


def _grouped(pairs: list[tuple[int, str]], title: str) -> list[dict]:
    """An outline from (level, title) pairs of levels 1 and 2, as markdown_outline builds one."""
    outline: list[dict] = []
    for level, heading in pairs:
        if level == 1:
            outline.append({"title": heading, "topics": []})
        else:
            if not outline:
                outline.append({"title": title, "topics": []})
            outline[-1]["topics"].append({"title": heading})
    return outline


def _pdf_outline(reader: PdfReader) -> list[tuple[int, str]]:
    """A PDF's first- and second-level bookmarks, as (level, title) pairs."""
    pairs = []
    for item in reader.outline:
        if isinstance(item, list):
            pairs.extend((2, plain(child.title)) for child in item if not isinstance(child, list))
        else:
            pairs.append((1, plain(item.title)))
    return pairs


def file_outline(path: Path) -> list[dict]:
    """A file's outline: Markdown's # and ## headings, Word's Heading 1 and Heading 2, a PDF's bookmarks.

    Raises:
        SourceError: When the file can't be read.
    """
    suffix = path.suffix.lower()
    if suffix in (".md", ".markdown"):
        return markdown_outline(read_file(path), title=path.stem)
    if suffix == ".pdf":
        return _grouped(_pdf_outline(_pdf(path)), path.stem)
    if path.stat().st_size > FOLDER_MAX_BYTES:
        raise SourceError("This file is too large to read.")
    try:
        paragraphs = Document(str(path)).paragraphs
    except (OSError, ValueError) as error:
        raise SourceError("This file could not be read.") from error
    levels = {"Heading 1": 1, "Heading 2": 2}
    return _grouped([(levels[paragraph.style.name], plain(paragraph.text)) for paragraph in paragraphs
                     if paragraph.style is not None and paragraph.style.name in levels and paragraph.text.strip()], path.stem)


def briefing_of(text: str) -> str | None:
    """What a Markdown text is about: its first paragraph of prose, cut to the word limit (see wording).

    Headings, code, lists, tables, quotes and images are passed over.

    Returns:
        The briefing, or None when the text has no paragraph of prose.
    """
    paragraph: list[str] = []
    for _, line, fenced in _lines(text):
        stripped = line.strip()
        prose = (not fenced and stripped and not _HEADING.match(stripped)
                 and not re.match(r"^([-*+>|]|\d+[.)]\s|!\[)", stripped))
        if prose:
            paragraph.append(stripped)
        elif paragraph:
            break
    return cut_words(plain(" ".join(paragraph))) or None


def file_briefing(path: Path) -> str | None:
    """What a file is about: Markdown's first paragraph, or the first paragraph of a Word file or a PDF's text."""
    text = read_file(path)
    if path.suffix.lower() in (".md", ".markdown"):
        return briefing_of(text)
    for block in re.split(r"\n\s*\n|\n(?=[A-Z])", text):
        if count_words(block) >= 5:
            return cut_words(plain(block)) or None
    return None


class _Page(HTMLParser):
    """Reads a page's title, its own description and its h1 and h2 headings, passing over everything else."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title, self.description, self.social, self.headings = "", None, None, []
        self._in, self._text, self._skip = None, [], 0

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if tag in ("script", "style", "template", "noscript"):
            self._skip += 1
        elif tag == "meta" and values.get("content"):
            if (values.get("name") or "").lower() == "description":
                self.description = self.description or values["content"].strip()
            elif (values.get("property") or "").lower() == "og:description":
                self.social = self.social or values["content"].strip()
        elif tag in ("title", "h1", "h2") and not self._skip:
            self._in, self._text = tag, []

    def handle_endtag(self, tag):
        if tag in ("script", "style", "template", "noscript"):
            self._skip = max(0, self._skip - 1)
        elif tag == self._in:
            text = " ".join("".join(self._text).split())
            if tag == "title":
                self.title = self.title or text
            elif text:
                self.headings.append((1 if tag == "h1" else 2, text))
            self._in = None

    def handle_data(self, data):
        if self._in and not self._skip:
            self._text.append(data)


def parse_page(html: str) -> dict:
    """What DayWright keeps of a web page: its title, its own briefing (its description, cut to the word limit), and
    its h1 headings each with its h2 ones. None of its other text is kept.

    Returns:
        {"title", "briefing", "outline": [{"title", "topics": [{"title"}]}]}; briefing is None when
        the page publishes none.
    """
    page = _Page()
    page.feed(html)
    page.close()
    title = page.title or None
    return {"title": title, "briefing": cut_words(page.description or page.social or "") or None,
            "outline": _grouped(page.headings, title or "")}


def fetch_page(address: str) -> dict:
    """Look a website up: read its page once, over http or https, within FETCH_TIMEOUT seconds and
    FETCH_MAX_BYTES, and keep what parse_page keeps.

    Raises:
        SourceError: For an address that isn't a web address, or a site that can't be reached.
    """
    parts = urlsplit(address)
    if parts.scheme not in ("http", "https") or not parts.netloc:
        raise SourceError("Give a web address that starts with http:// or https://.")
    request = Request(address, headers={"User-Agent": "DayWright", "Accept": "text/html"})
    try:
        with urlopen(request, timeout=FETCH_TIMEOUT) as answer:
            data = answer.read(FETCH_MAX_BYTES)
            kind = answer.headers.get("Content-Type", "") or ""
    except (OSError, ValueError) as error:
        raise SourceError("The website could not be reached.") from error
    charset = re.search(r"charset=([\w-]+)", kind)
    return parse_page(data.decode(charset.group(1) if charset else "utf-8", errors="replace"))


def slugify(value: str) -> str:
    """A lesson's slug as Observatory makes it from its path: lower case, & as "and", every other run
    of characters that aren't letters or digits as one dash."""
    text = _BLOCK_ANCHOR.sub("", unquote(value)).lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def lesson_address(site: str, relative: str) -> str:
    """Where a folder's file opens on the website the user gave the folder.

    When that address is the Observatory site, a Markdown file is a lesson with its own page and any
    other file opens its top folder's collection page; any other address opens at its main page.

    Args:
        site: The folder's website, as the user entered it.
        relative: The file's path in the folder, such as "Guides/01 Basics.md".
    """
    if urlsplit(site).netloc not in OBSERVATORY_HOSTS:
        return site
    root = site.rstrip("/")
    path = PurePosixPath(relative)
    if path.suffix.lower() in (".md", ".markdown"):
        return f"{root}/lesson/{slugify(str(path.with_suffix('')))}"
    return f"{root}/library/{slugify(path.parts[0])}" if len(path.parts) > 1 else root


def checklist_of(outline: list[dict]) -> list[dict]:
    """The checklist a task made from a source holds: its second-level headings, in order; with none, its
    first-level ones below the first, which names it; with none of those, nothing.

    Returns:
        [{"title", "level"}], level being the headings' own, 2 or 1.
    """
    sections = [{"title": topic["title"], "level": 2} for group in outline for topic in group["topics"]]
    return sections or [{"title": group["title"], "level": 1} for group in outline[1:]]


def section_text(text: str, heading: str, level: int = 2) -> str:
    """The text under a heading of the second level, or of `level`, to the next heading of that level or a higher one."""
    lines = text.splitlines()
    start = end = None
    for number, line, fenced in _lines(text):
        match = None if fenced else _HEADING.match(line)
        if not match or len(match.group(1)) > level:
            continue
        if start is None and len(match.group(1)) == level and plain(match.group(2)) == heading:
            start = number + 1
        elif start is not None:
            end = number
            break
    return "" if start is None else "\n".join(lines[start:end]).strip()


def study_profile(text: str) -> dict:
    """What a text says about studying it, a whole file's or the sections left of it, worked out the same
    way every time.

    Returns:
        "subheadings" (its third-level headings: what will be learnt), "words" (outside code),
        "codeBlocks", "handsOn" (code or an exercise in it), and "effort": light under LIGHT_WORDS words
        and not hands-on, deep from DEEP_WORDS words or DEEP_CODE_BLOCKS code examples, else steady.
    """
    subheadings, prose, fences = [], [], 0
    for _, line, fenced in _lines(text):
        if _FENCE.match(line):
            fences += 1
            continue
        if fenced:
            continue
        match = _HEADING.match(line)
        if match:
            if len(match.group(1)) == 3:
                subheadings.append(plain(match.group(2)))
            continue
        prose.append(line)
    words_text = plain("\n".join(prose))
    words = len(_WORD.findall(words_text)) + len(_CJK.findall(words_text)) // 2
    code_blocks = fences // 2
    hands_on = bool(code_blocks or _HANDS_ON.search(words_text))
    return {"subheadings": subheadings, "words": words, "codeBlocks": code_blocks, "handsOn": hands_on,
            "effort": _effort(words, code_blocks, hands_on)}


def _effort(words: int, code_blocks: int, hands_on: bool) -> str:
    """Light under LIGHT_WORDS words and not hands-on, deep from DEEP_WORDS words or DEEP_CODE_BLOCKS code
    examples, else steady."""
    return ("deep" if words >= DEEP_WORDS or code_blocks >= DEEP_CODE_BLOCKS
            else "light" if words < LIGHT_WORDS and not hands_on else "steady")


def combined_profile(profiles: list[dict]) -> dict:
    """The profile of several parts studied as one, such as a file's sections read one by one: their words,
    code examples and subheadings together, and the effort that gives; parts with no text, a website's,
    are weighed by their count (see heading_profile)."""
    if not any("words" in profile for profile in profiles):
        return heading_profile(len(profiles))
    words = sum(profile.get("words", 0) for profile in profiles)
    code_blocks = sum(profile.get("codeBlocks", 0) for profile in profiles)
    hands_on = any(profile.get("handsOn") for profile in profiles)
    return {"subheadings": [heading for profile in profiles for heading in profile.get("subheadings", [])],
            "words": words, "codeBlocks": code_blocks, "handsOn": hands_on, "effort": _effort(words, code_blocks, hands_on)}


def heading_profile(sections: int) -> dict:
    """What a website's page says about studying it, from its headings alone, as it keeps no text.

    Returns:
        "sections" (how many it covers) and "effort": light up to LIGHT_SECTIONS, deep from DEEP_SECTIONS,
        else steady.
    """
    effort = "light" if sections <= LIGHT_SECTIONS else "deep" if sections >= DEEP_SECTIONS else "steady"
    return {"sections": sections, "effort": effort}


def study_minutes(profile: dict) -> int:
    """How long a study session takes, in whole slots and never under MIN_STUDY_MINUTES: a text's words at
    STUDY_WORDS_PER_MINUTE and CODE_BLOCK_MINUTES for each code example, and SECTION_MINUTES for each section
    with no text of its own, such as a website's (see heading_profile) or a checklist item the user added."""
    minutes = (profile.get("words", 0) / STUDY_WORDS_PER_MINUTE + profile.get("codeBlocks", 0) * CODE_BLOCK_MINUTES
               + profile.get("sections", 0) * SECTION_MINUTES)
    return max(MIN_STUDY_MINUTES, ceil(minutes / SLOT_MINUTES) * SLOT_MINUTES)
